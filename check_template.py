import openpyxl
from docx import Document

# Check Excel full column names
print('=== Excel Columns ===')
wb = openpyxl.load_workbook('d:/03coding_tools/vscode/01ETW/ETW/01安全审查记录-new.xlsx')
ws = wb.active
headers = [str(ws.cell(1, ci).value or '') for ci in range(1, ws.max_column+1)]
for i, h in enumerate(headers):
    print(f'  Col {i+1}: {h}')

# Check Excel row 2 data
print('\n=== Excel Row 2 Data ===')
for ci in range(1, ws.max_column+1):
    h = ws.cell(1, ci).value
    v = ws.cell(2, ci).value
    print(f'  {h}: {v}')

# Check Word Table structure
print('\n=== Word Table 1 Structure ===')
doc = Document('d:/03coding_tools/vscode/01ETW/ETW/01安全审查记录-new.docx')
t = doc.tables[0]
print(f'Table 1: {len(t.rows)} rows')
for ri, row in enumerate(t.rows):
    col1 = row.cells[0].text.strip()
    col2 = row.cells[1].text.strip()[:50]
    print(f'  Row {ri}: [{col1}] = [{col2}]')